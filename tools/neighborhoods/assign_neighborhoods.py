#!/usr/bin/env python3
"""Assign one Jerusalem neighborhood to every address in an Excel list.

Input:  an .xlsx whose first column holds addresses ("street number"), with or
        without a header row.
Output: an .xlsx with exactly two columns, רחוב (the original address) and
        שכונה (one name from the agent's neighborhood registry, or blank),
        sorted neighborhood by neighborhood in the registry's order.

How each address is decided, first match wins:
  1. data/overrides.json "addresses"  — a fix made by hand for one address.
  2. data/address_cache.csv           — the decision made in an earlier run.
  3. data/overrides.json "streets"    — a fix made by hand for a whole street
                                        ("A / B" narrows it to those options).
  4. data/municipal_streets_2020.xlsx — the municipality's street → neighborhood
     list (exact name, then spelling-normalized, then a close spelling).
  5. A street listed under several neighborhoods is located on OpenStreetMap:
     the neighborhood OSM records at that point if it is one of the street's
     options, else the option whose center is nearest. Not found on the map →
     the option with the most rows in the municipal list.

New decisions are added to the caches, so the next run only works on new
addresses. See README.md.

Usage:
  python3 assign_neighborhoods.py NEW_ADDRESSES.xlsx OUT.xlsx [--no-cache] [--no-geocode]
"""
import argparse
import collections
import csv
import difflib
import json
import math
import re
import sys
import time
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
DATA = HERE / 'data'
REPO_DATA = HERE.parents[1] / 'data'   # Arnona_Agent keeps the registry in <repo>/data
RAW = 'https://raw.githubusercontent.com/urbandetective007/Arnona_Agent/main/data/'
USER_AGENT = 'ArnonaAgentProject/1.0 (Jerusalem municipality property map)'
BOUNDS = dict(south=31.70, north=31.90, west=35.07, east=35.32)


# ── Registry (closed neighborhood list + map centers) ─────────────────────────

def load_registry_file(name):
    """From <repo>/data when run inside Arnona_Agent, else from Arnona_Agent on GitHub."""
    local = REPO_DATA / name
    if local.exists() and (REPO_DATA / 'jerusalem_neighborhood_centers.json').exists():
        return json.loads(local.read_text(encoding='utf-8'))
    req = urllib.request.Request(RAW + name, headers={'User-Agent': USER_AGENT})
    with urllib.request.urlopen(req, timeout=30) as r:
        return json.load(r)


# ── Text normalization ────────────────────────────────────────────────────────

def clean(s):
    return ' '.join(str(s).split())


def norm(s):
    s = re.sub(r"[\"'`׳״’]", '', str(s)).replace('-', ' ')
    return re.sub(r'\s+', ' ', s).strip()


def drop_alley(s):
    return re.sub(r'\s+סמ\s*\d+$', '', s).strip()


def drop_prefix(s):
    for p in ['שכ ', 'שכונת ', 'דרך ', 'רחוב ', 'שדרות ', 'שד ', 'הרבנית ', 'הרב ', 'האדמור ', 'רבי ', 'ר ']:
        if s.startswith(p):
            s = s[len(p):]
    return s


def sorted_words(s):
    return ' '.join(sorted(s.split()))


def street_of(address):
    """'דרך שכם סמ5 12' → 'דרך שכם סמ5' (drops the trailing house number)."""
    return re.sub(r'\s+\S*\d\S*$', '', clean(address)).strip()


# ── Municipal street list ─────────────────────────────────────────────────────

class StreetIndex:
    LEVELS = ('exact', 'norm', 'alley', 'prefix', 'words')
    NOT_CLOSE = {'אל סילם'}   # close to 'אל סיל' but a different street

    def __init__(self, path):
        self.idx = {k: {} for k in self.LEVELS}
        ws = openpyxl.load_workbook(path, read_only=True).active
        for street, hood in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            if not street or not hood:
                continue
            for level, key in zip(self.LEVELS, self.keys(str(street).strip())):
                self.idx[level].setdefault(key, collections.Counter())[str(hood).strip()] += 1
        self.fuzzy_keys = list(self.idx['prefix'])

    @staticmethod
    def keys(street):
        n = norm(street)
        p = drop_prefix(drop_alley(n))
        return street.strip(), n, drop_alley(n), p, sorted_words(p)

    def lookup(self, street):
        """Municipal neighborhood names for a street (most rows first), and how they were found."""
        for level, key in zip(self.LEVELS, self.keys(street)):
            if key in self.idx[level]:
                return self.idx[level][key], 'exact' if level == 'exact' else 'spelling'
        n = norm(street)
        if n.startswith('שכ '):   # "שכ <name>" names the neighborhood itself
            return collections.Counter({re.sub(r'\s+גוש.*$', '', street.strip()[3:]).strip(): 1}), 'named'
        key = drop_prefix(drop_alley(n))
        # Close spellings only for longer names — short ones matched wrong streets (צורי → צפורי).
        close = [] if len(key) < 7 or key in self.NOT_CLOSE else difflib.get_close_matches(key, self.fuzzy_keys, 1, 0.88)
        if close:
            return self.idx['prefix'][close[0]], 'close'
        return None, 'none'


# ── Neighborhood names → registry names ───────────────────────────────────────

def name_key(s):
    s = re.sub(r"[\"'`׳״’\-\s]", '', s)
    return re.sub(r'וו', 'ו', s).replace('שייח', 'שיח')


def loose_key(s):
    return re.sub(r'^ה', '', name_key(re.sub(r'(^|\s)אל[\s\-]', ' ', s)))


class Names:
    def __init__(self, registry, name_map):
        self.canon = list(registry['neighborhoods'])
        self.canon_set = set(self.canon)
        self.map = {**registry['aliases'], **name_map}
        self.by_key = {name_key(n): n for n in self.canon}
        self.by_loose = {loose_key(n): n for n in self.canon}
        self.unknown = collections.Counter()

    def to_registry(self, raw, count_unknown=True):
        """A municipal / OSM name → registry names (a combined area can give two)."""
        raw = raw.strip()
        mapped = self.map.get(raw, raw)
        out = []
        for part in mapped.split(' / '):
            part = part.strip()
            hit = part if part in self.canon_set else (
                self.by_key.get(name_key(part)) or self.by_loose.get(loose_key(part)))
            if hit:
                out.append(hit)
            elif count_unknown:
                self.unknown[part] += 1
        return out


# ── OpenStreetMap ─────────────────────────────────────────────────────────────

def km(a, b):
    return math.hypot((a['lat'] - b['lat']) * 111.32,
                      (a['lon'] - b['lon']) * 111.32 * math.cos(math.radians(a['lat'])))


def address_queries(address):
    """Same spellings Arnona_Agent's src/lib/geocode.ts tries."""
    m = re.match(r'^(.*?\D)\s*(\d+)(?:\s*[א-ת]|\s*/\s*[\dא-ת]+)?(?:\s+.*)?$', clean(address))
    if not m or m.group(2) == '0':
        return []
    street, number = m.group(1).strip(), m.group(2)
    variants = [street]
    if street.startswith('הרב '):
        variants.append(street[4:])
    if '"' in street:
        variants += [street.replace('"', '״'), street.replace('"', '')]
    words = street.split(' ')
    not_person = {'בית', 'גבעת', 'דרך', 'עין', 'הר', 'שדה', 'נוף', 'רמת', 'קרית', 'כפר', 'נווה', 'מבוא'}
    if len(words) == 2 and not words[0].startswith('ה') and words[0] not in not_person:
        variants.append(f'{words[1]} {words[0]}')
    return list(dict.fromkeys([clean(address)] + [f'{v} {number}' for v in variants]))


def nominatim(query):
    params = urllib.parse.urlencode(dict(
        format='json', limit=5, addressdetails=1, countrycodes='il', bounded=1,
        viewbox=f"{BOUNDS['west']},{BOUNDS['north']},{BOUNDS['east']},{BOUNDS['south']}",
        q=f'{query}, ירושלים'))
    req = urllib.request.Request('https://nominatim.openstreetmap.org/search?' + params,
                                 headers={'User-Agent': USER_AGENT, 'Accept-Language': 'he'})
    wait = 30
    for attempt in range(6):
        time.sleep(1.3)   # Nominatim allows ~1 request per second
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                return json.load(r)
        except Exception as e:  # rate limit / network: back off and retry
            print(f'  retry ({e}) in {wait}s', file=sys.stderr)
            time.sleep(wait)
            wait = min(wait * 2, 300)
    raise RuntimeError(f'Nominatim unreachable for {query!r}')


def in_jerusalem(r):
    a = r.get('address', {})
    city = next((a[k] for k in ('city', 'town', 'municipality', 'village') if a.get(k)), '')
    lat, lon = float(r['lat']), float(r['lon'])
    return (re.search(r'ירושלים|jerusalem', city, re.I) and
            BOUNDS['south'] <= lat <= BOUNDS['north'] and BOUNDS['west'] <= lon <= BOUNDS['east'])


def locate(address, near):
    """Building-level hit if possible, else street-level, else not found."""
    def pick(ms):
        if not near or len(ms) <= 1:
            return ms[0]
        return min(ms, key=lambda m: km({'lat': float(m['lat']), 'lon': float(m['lon'])}, near))
    street_hit = None
    for q in address_queries(address):
        ms = [m for m in nominatim(q) if in_jerusalem(m)]
        if not ms:
            continue
        exact = [m for m in ms if m.get('address', {}).get('house_number')]
        if exact:
            hit, precision = pick(exact), 'exact'
            break
        street_hit = street_hit or pick(ms)
    else:
        if not street_hit:
            return {'found': False}
        hit, precision = street_hit, 'street'
    a = hit.get('address', {})
    return {'found': True, 'lat': float(hit['lat']), 'lon': float(hit['lon']), 'precision': precision,
            'osm_area': [x for x in (a.get('neighbourhood'), a.get('quarter'), a.get('suburb')) if x]}


def choose(options, geo, names, centers):
    """One of `options` (registry names, most municipal rows first) for a located address."""
    if not geo.get('found'):
        return options[0], 'not on map → most common'
    for area in geo['osm_area']:
        for n in names.to_registry(area, count_unknown=False):
            if n in options:
                return n, f'OSM area ({geo["precision"]})'
    p = {'lat': geo['lat'], 'lon': geo['lon']}
    with_center = [o for o in options if o in centers]
    if not with_center:
        return options[0], 'no centers → most common'
    return min(with_center, key=lambda o: km(p, centers[o])), f'nearest center ({geo["precision"]})'


# ── Caches ────────────────────────────────────────────────────────────────────

def read_address_cache():
    path = DATA / 'address_cache.csv'
    if not path.exists():
        return {}
    with path.open(encoding='utf-8', newline='') as f:
        return {row['address']: row['neighborhood'] for row in csv.DictReader(f) if row['neighborhood']}


def write_address_cache(cache):
    with (DATA / 'address_cache.csv').open('w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        w.writerow(['address', 'neighborhood'])
        for a in sorted(cache):
            w.writerow([a, cache[a]])


def read_geo_cache():
    path = DATA / 'geocode_cache.jsonl'
    if not path.exists():
        return {}
    return {r['address']: r for r in map(json.loads, path.read_text(encoding='utf-8').splitlines()) if r}


# ── Main ──────────────────────────────────────────────────────────────────────

def read_addresses(path):
    ws = openpyxl.load_workbook(path, read_only=True).active
    rows = [r[0] for r in ws.iter_rows(max_col=1, values_only=True)]
    if rows and str(rows[0]).strip() in ('רחוב', 'כתובת', 'address'):
        rows = rows[1:]
    return [r for r in rows if r is not None and str(r).strip()]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('input')
    ap.add_argument('output')
    ap.add_argument('--no-cache', action='store_true', help='ignore earlier decisions (they are still updated)')
    ap.add_argument('--no-geocode', action='store_true',
                    help="don't call OpenStreetMap; streets in several neighborhoods get the most common one")
    args = ap.parse_args()

    registry = load_registry_file('jerusalem_neighborhoods.json')
    centers = load_registry_file('jerusalem_neighborhood_centers.json')
    overrides = json.loads((DATA / 'overrides.json').read_text(encoding='utf-8'))
    names = Names(registry, overrides.get('name_map', {}))
    streets = StreetIndex(DATA / 'municipal_streets_2020.xlsx')
    addr_cache = read_address_cache()
    geo_cache = read_geo_cache()
    geo_log = (DATA / 'geocode_cache.jsonl').open('a', encoding='utf-8')
    addr_over = {clean(k): v for k, v in overrides.get('addresses', {}).items()}
    street_over = {norm(k): v for k, v in overrides.get('streets', {}).items()}

    addresses = read_addresses(args.input)
    print(f'{len(addresses)} addresses', file=sys.stderr)
    stats, missing = collections.Counter(), collections.Counter()
    out = openpyxl.Workbook()
    ws = out.active
    ws.sheet_view.rightToLeft = True
    ws.append(['רחוב', 'שכונה'])
    results = []
    for i, raw in enumerate(addresses, 1):
        a = clean(raw)
        street = street_of(a)
        hood, how = '', ''
        if a in addr_over:
            hood, how = addr_over[a], 'manual (address)'
        elif not args.no_cache and a in addr_cache:
            hood, how = addr_cache[a], 'earlier run'
        elif norm(street) in street_over and ' / ' not in street_over[norm(street)]:
            hood, how = street_over[norm(street)], 'manual (street)'
        else:
            if norm(street) in street_over:   # 'A / B': a hand-narrowed list of options
                found, match = None, 'manual'
                options = street_over[norm(street)].split(' / ')
            else:
                found, match = streets.lookup(street)
                options = []
                for n, _ in (found.most_common() if found else []):
                    options += [x for x in names.to_registry(n) if x not in options]
            if not options:
                how = 'street not in municipal list' if not found else 'name not in registry'
                missing[street] += 1
            elif len(options) == 1:
                hood, how = options[0], f'municipal list ({match})'
            elif args.no_geocode:
                hood, how = options[0], 'several → most common (no geocode)'
            else:
                if a not in geo_cache:
                    pts = [centers[o] for o in options if o in centers]
                    near = {'lat': sum(p['lat'] for p in pts) / len(pts),
                            'lon': sum(p['lon'] for p in pts) / len(pts)} if pts else None
                    geo_cache[a] = {'address': a, **locate(a, near)}
                    geo_log.write(json.dumps(geo_cache[a], ensure_ascii=False) + '\n')
                    geo_log.flush()
                hood, how = choose(options, geo_cache[a], names, centers)
                how = 'several → ' + how
        if hood and hood not in names.canon_set:
            print(f'  ! "{hood}" for {a} is not in the registry — left blank', file=sys.stderr)
            hood, how = '', 'not in registry'
        if hood:
            addr_cache[a] = hood
        stats[how] += 1
        results.append((raw, hood))
        if i % 1000 == 0:
            print(f'  {i}/{len(addresses)}', file=sys.stderr)
    geo_log.close()
    # The agent scans rows in order, so group them neighborhood by neighborhood, in the
    # registry's order (nearby neighborhoods are listed together); street order is kept
    # inside each neighborhood, and addresses without a neighborhood come last.
    rank = {n: i for i, n in enumerate(names.canon)}
    for row in sorted(results, key=lambda r: (rank[r[1]], 0) if r[1] in rank else (len(rank), 1 if r[1] else 2)):
        ws.append(list(row))
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 25
    out.save(args.output)
    write_address_cache(addr_cache)

    print('\nHow addresses were decided:')
    for how, n in stats.most_common():
        print(f'  {n:6}  {how}')
    if missing:
        print(f'\nStreets not in the municipal list ({len(missing)}), most common:')
        for s, n in missing.most_common(20):
            print(f'  {n:4}  {s}')
    if names.unknown:
        print('\nMunicipal names with no registry match (add to overrides.json "name_map"):')
        for s, n in names.unknown.most_common(20):
            print(f'  {n:4}  {s}')


if __name__ == '__main__':
    main()
