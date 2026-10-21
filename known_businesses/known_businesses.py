#!/usr/bin/env python3
"""עסקים מוכרים — בודק אם עסק כבר נמצא בעבר (במערכת או בריצה קודמת של הרוטינה).

  check "<שם העסק>" "<כתובת>"        → KNOWN (כבר קיים, לא להכניס לדוח) / NEW (חדש)
  check-report "<דוח>.xlsx"           → בודק כל שורה בדוח ומדפיס KNOWN/NEW לכל אחת
  add "<דוח>.xlsx" --source "<מקור>"  → מוסיף לקובץ את העסקים מהדוח (רק מי שלא קיים כבר)

ההתאמה: אותה כתובת (רחוב + מספר בית) וגם שם דומה — אחרי ניקוי תארים (עו"ד, ד"ר, בע"מ...),
סימני פיסוק וסדר מילים. כך "בן צבי נרדה" ו-"נרדה בן-צבי, עו\"ד" באותה כתובת נחשבים אותו עסק.
"""
import argparse
import os
import re
import sys
from datetime import date

import openpyxl
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
KNOWN_FILE = os.path.join(HERE, "known_businesses.xlsx")
HEADERS = ["שם העסק", "כתובת", "שכונה", "סוג עסק", "קישור", "מקור", "תאריך הוספה"]

# מילים שאינן חלק מזהה של השם
NAME_STOPWORDS = {
    "עוד", "עו״ד", "עורך", "עורכת", "עורכי", "דין", "ונוטריון", "נוטריון", "משרד", "משרדי",
    "דר", "ד״ר", "דוקטור", "פרופ", "פרופסור", "מר", "גב", "הרב", "רב",
    "בעמ", "בע״מ", "חברה", "ושות", "ושותפים", "שותפים", "קבוצת", "ירושלים",
    "the", "ltd", "inc", "co", "dr", "adv",
}
ADDRESS_DROP = {"ירושלים", "jerusalem", "רחוב", "רח", "ישראל", "israel"}
STREET_PREFIX = {"שדרות": "שד", "שדרת": "שד", "שד": "שד", "דרך": "דרך", "סמטת": "סמ", "סמ": "סמ", "כיכר": "ככר", "ככר": "ככר"}


def _clean(text):
    t = str(text or "").lower()
    t = t.replace('"', "").replace("״", "").replace("'", "").replace("׳", "").replace("`", "")
    t = re.sub(r"[-–—_.,;:()\[\]/\\|+]", " ", t)
    return t.split()


def address_key(address):
    """('<רחוב מנורמל>', '<מספר בית>') — או None אם אין מספר בית."""
    raw = str(address or "")
    raw = re.sub(r"\b\d{7}\b", " ", raw)  # מיקוד
    raw = re.sub(r"(\d+)\s*/\s*\d+", r"\1", raw)  # 1/2 → 1
    raw = re.sub(r"(\d+)\s*[אבגדה]\b", r"\1", raw)  # 12א → 12
    tokens = [t for t in _clean(raw) if t not in ADDRESS_DROP]
    numbers = [t for t in tokens if t.isdigit()]
    if not numbers:
        return None
    words = [STREET_PREFIX.get(t, t) for t in tokens if not t.isdigit()]
    return (" ".join(sorted(words)), numbers[0])


ARABIC_STOPWORDS = {"الدكتور", "دكتور", "د", "المحامي", "محامي", "مكتب", "شركة", "القدس", "ال"}
ARABIC_LETTERS = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ى": "ي", "ة": "ه", "ـ": ""})
ARABIC_DIACRITICS = re.compile(r"[ً-ْٰ]")


def _norm_token(t):
    """מאחד תעתיקים שונים של אותו שם ערבי: "מוחמד" = "מחמד", "אלחוסייני" = "חוסייני", "أحمد" = "احمد"."""
    if re.search(r"[؀-ۿ]", t):
        t = ARABIC_DIACRITICS.sub("", t).translate(ARABIC_LETTERS)
        return t[2:] if t.startswith("ال") and len(t) > 4 else t
    if t.startswith("אל") and len(t) >= 5:
        t = t[2:]
    return t[0] + re.sub(r"[וי]", "", t[1:]) if len(t) > 2 else t


def name_tokens(name):
    tokens = (t for t in _clean(name) if t not in NAME_STOPWORDS and t not in ARABIC_STOPWORDS and t not in {"אל", "א"})
    return {_norm_token(t) for t in tokens if len(t) > 1}


def names_match(a, b):
    ta, tb = name_tokens(a), name_tokens(b)
    if not ta or not tb:
        return False
    if ta <= tb or tb <= ta:
        return True
    return len(ta & tb) / len(ta | tb) >= 0.5


def load_known():
    if not os.path.exists(KNOWN_FILE):
        return []
    ws = openpyxl.load_workbook(KNOWN_FILE, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    idx = {h: i for i, h in enumerate(rows[0]) if h}
    out = []
    for r in rows[1:]:
        rec = {h: (r[i] if i < len(r) else None) for h, i in idx.items()}
        if rec.get("שם העסק"):
            out.append(rec)
    return out


def find_match(name, address, known):
    """מחזיר (רשומה תואמת, 'same_address' / 'name_only') או (None, None)."""
    key = address_key(address)
    name_only = None
    for rec in known:
        rkey = address_key(rec.get("כתובת"))
        if key and rkey and key == rkey and names_match(name, rec["שם העסק"]):
            return rec, "same_address"
        if name_only is None and name_tokens(name) and name_tokens(name) == name_tokens(rec["שם העסק"]) and len(name_tokens(name)) >= 2:
            name_only = rec
    if name_only:
        return name_only, "name_only"
    return None, None


def save_known(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "עסקים מוכרים"
    ws.sheet_view.rightToLeft = True
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    for r in rows:
        ws.append([r.get(h) for h in HEADERS])
    for col, width in zip("ABCDEFG", (38, 32, 18, 26, 45, 28, 13)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    wb.save(KNOWN_FILE)


def read_report(path):
    ws = openpyxl.load_workbook(path, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    idx = {str(h).strip(): i for i, h in enumerate(rows[0]) if h}
    if "שם העסק" not in idx or "כתובת" not in idx:
        sys.exit(f"❌ {path}: אין עמודות 'שם העסק' ו-'כתובת' בשורה הראשונה")
    get = lambda r, h: (str(r[idx[h]]).strip() if h in idx and idx[h] < len(r) and r[idx[h]] is not None else None)
    out = []
    for r in rows[1:]:
        name, addr = get(r, "שם העסק"), get(r, "כתובת")
        if name and addr:
            link = next((get(r, h) for h in ("קישור 1", "קישור", "קישור 2", "קישור 3") if get(r, h)), None)
            out.append({"שם העסק": name, "כתובת": addr, "שכונה": get(r, "שכונה"),
                        "סוג עסק": get(r, "סוג עסק") or get(r, "סוג העסק"), "קישור": link})
    return out


def report_line(name, address, rec, kind):
    if kind == "same_address":
        return f"KNOWN | {name} | {address} | כבר קיים: {rec['שם העסק']} | {rec['כתובת']} | מקור: {rec.get('מקור')}"
    if kind == "name_only":
        return (f"NEW (שם זהה בכתובת אחרת — לבדוק) | {name} | {address} | "
                f"קיים: {rec['שם העסק']} | {rec['כתובת']} | מקור: {rec.get('מקור')}")
    return f"NEW | {name} | {address}"


def main():
    p = argparse.ArgumentParser()
    sub = p.add_subparsers(dest="cmd", required=True)
    c = sub.add_parser("check")
    c.add_argument("name")
    c.add_argument("address")
    cr = sub.add_parser("check-report")
    cr.add_argument("report")
    a = sub.add_parser("add")
    a.add_argument("report")
    a.add_argument("--source", required=True)
    args = p.parse_args()

    known = load_known()
    if args.cmd == "check":
        rec, kind = find_match(args.name, args.address, known)
        print(report_line(args.name, args.address, rec, kind))
        sys.exit(1 if kind == "same_address" else 0)
    if args.cmd == "check-report":
        n_known = 0
        for b in read_report(args.report):
            rec, kind = find_match(b["שם העסק"], b["כתובת"], known)
            n_known += kind == "same_address"
            print(report_line(b["שם העסק"], b["כתובת"], rec, kind))
        print(f"סה\"כ KNOWN: {n_known}")
        sys.exit(1 if n_known else 0)
    if args.cmd == "add":
        added = 0
        for b in read_report(args.report):
            if find_match(b["שם העסק"], b["כתובת"], known)[1] == "same_address":
                continue
            b["מקור"] = args.source
            b["תאריך הוספה"] = date.today().isoformat()
            known.append(b)
            added += 1
        save_known(known)
        print(f"✅ נוספו {added} עסקים. סה\"כ בקובץ: {len(known)}")


if __name__ == "__main__":
    main()
